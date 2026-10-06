#!/usr/bin/env python3
"""技術科「生物育成」栽培観察レポートを、生徒が Word で記入できる形式（.docx）で作成する。

構成（すべて A4 縦・1日1枚）:
  1枚目      表紙（作物名入りタイトル・年組番名前・観察・栽培期間）
  2枚目〜    月〜金の観察日記（写真・一番大きい株の背丈・葉の枚数・手入れ・考察）
  最後の1枚  収穫後のまとめ

記入欄はすべて Word の「コンテンツ コントロール」:
  - 文字欄：クリックして入力（灰色の案内文字は入力すると消える）
  - 天気：一覧から選ぶ　／　手入れ：クリックで ☐→☑　／　収穫日：カレンダーから選ぶ
  - 写真欄：クリックすると「図の挿入」が開き、写真が枠の大きさに合わせて入る
  - 表紙で名前などを入れると、全ページのヘッダーにも自動で反映される
--print を付けると、手書き用（罫線つき・コントロールなし）で作る。

例:
  python3 make_report.py --crop ミニトマト --start 2026-05-11 --end 2026-07-17
  python3 make_report.py --crop ダイコン --start 2026-10-05 --end 2026-12-18 \
      --grade 2 --klass 3 --number 15 --name "山田 太郎" \
      --skip 2026-11-03,2026-11-23 --photos ./photos --pdf -o daikon.docx
"""
import argparse
import datetime as dt
import io
import os
import re
import subprocess
import sys
import uuid
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Mm, Pt, RGBColor

FONT = "BIZ UDPGothic"          # Windows 10 以降に標準搭載。無い環境では代替フォントで表示される
BOX_FONT = "MS Gothic"          # チェックボックス（☐☑）用
WEEKDAY = "月火水木金土日"
PAGE_W, PAGE_H = 210, 297       # A4 (mm)
MARGIN_X, MARGIN_TOP, MARGIN_BOTTOM = 15, 17, 14
BODY_W = PAGE_W - MARGIN_X * 2  # 180mm
GRAY = RGBColor(0x8A, 0x94, 0xA3)
GRAY_HEX = "8A94A3"
ACCENT = "2F7D4F"               # 見出しの帯（緑）
PALE = "EAF4EC"
CARE_ITEMS = ["水やり", "追肥", "間引き", "除草", "土寄せ", "支柱・誘引", "芽かき・摘心"]
WEATHER = ["晴れ", "くもり", "雨", "雪", "晴れ時々くもり", "くもり時々雨"]
PHOTO_EXT = (".jpg", ".jpeg", ".png", ".gif", ".bmp")
XML_NS = "urn:saibai-report"

PRINT = False                   # --print（手書き用）なら True
STORE_ID = "{" + str(uuid.uuid4()).upper() + "}"
_ids = iter(range(1000, 10 ** 6))


# ---------- 低レベルの書式ヘルパー ----------

def set_cell_borders(cell, **edges):
    """edges: top/bottom/left/right = (val, size_eighths_pt, color) or None（線なし）"""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        if edge not in edges:
            continue
        spec = edges[edge]
        el = OxmlElement(f"w:{edge}")
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            val, sz, color = spec
            el.set(qn("w:val"), val)
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
        borders.append(el)


def box_borders(cell, color="8A94A3", sz=8):
    line = ("single", sz, color)
    set_cell_borders(cell, top=line, left=line, bottom=line, right=line)


def shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def set_cell_margins(table, top=0.8, bottom=0.8, left=1.8, right=1.8):
    mar = OxmlElement("w:tblCellMar")
    for edge, mm in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(int(mm * 56.7)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    table._tbl.tblPr.append(mar)


def fix_layout(table, widths_mm):
    """列幅を固定（Word が自動調整で崩さないように）"""
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for col, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths_mm):
        col.set(qn("w:w"), str(int(w * 56.7)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths_mm):
            cell.width = Mm(w)


def row_height(row, mm):
    row.height = Mm(mm)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY


def tight(par, before=0, after=0):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.0


def first_par(cell, align=None):
    par = cell.paragraphs[0]
    tight(par)
    if align is not None:
        par.alignment = align
    return par


def text(par, s, size=10.5, bold=False, color=None):
    run = par.add_run(s)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    return run


def write(cell, s, size=10.5, bold=False, color=None, align=None):
    par = first_par(cell, align)
    text(par, s, size, bold, color)
    return par


def spacer(doc, pt=4):
    p = doc.add_paragraph()
    tight(p)
    p.paragraph_format.line_spacing = Pt(pt)
    text(p, "", size=1)
    return p


def add_page_field(par):
    run = par.add_run()
    for tag, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "end")):
        el = OxmlElement(tag)
        if attr:
            el.set(qn("w:fldCharType"), attr)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = "PAGE"
        run._r.append(el)
    run.font.size = Pt(8)
    run.font.color.rgb = GRAY


# 要素の並び順（OOXML スキーマ順。順番が違うと Word が「読み取れない内容」と言うことがある）
_ORDER = {
    "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
             "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"],
    "tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
              "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
              "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"],
}


def normalize(root):
    for tag, order in _ORDER.items():
        rank = {qn("w:" + n): i for i, n in enumerate(order)}
        for el in root.iter(qn("w:" + tag)):
            kids = sorted(el, key=lambda k: rank.get(k.tag, 99))
            for k in kids:
                el.append(k)


# ---------- コンテンツ コントロール（記入欄） ----------

def _rpr(size, color=None, font=None):
    c = f'<w:color w:val="{color}"/>' if color else ""
    f = (f'<w:rFonts w:ascii="{font}" w:eastAsia="{font}" w:hAnsi="{font}" w:hint="eastAsia"/>'
         if font else "")
    return f"<w:rPr>{f}{c}<w:sz w:val=\"{int(size * 2)}\"/><w:szCs w:val=\"{int(size * 2)}\"/></w:rPr>"


def _run(s, size, color=None, font=None):
    return f'<w:r>{_rpr(size, color, font)}<w:t xml:space="preserve">{escape(s)}</w:t></w:r>'


def _bind(key):
    return (f"<w:dataBinding w:prefixMappings=\"xmlns:ns0='{XML_NS}'\" "
            f'w:xpath="/ns0:report[1]/ns0:{key}[1]" w:storeItemID="{STORE_ID}"/>')


def field(par, hint, size=11, alias="", value="", bind=None, kind="<w:text/>"):
    """1行の入力欄。印刷用では空白。bind を指定すると表紙とヘッダーで内容が連動する。"""
    if PRINT:
        text(par, value or "　" * max(2, len(hint)), size)
        return
    shown = value or hint
    plc = "" if value else "<w:showingPlcHdr/>"
    xml = (f"<w:sdt {nsdecls('w')}><w:sdtPr>{_rpr(size)}<w:alias w:val=\"{escape(alias or hint)}\"/>"
           f'<w:id w:val="{next(_ids)}"/>{plc}{_bind(bind) if bind else ""}{kind}</w:sdtPr>'
           f"<w:sdtContent>{_run(shown, size, None if value else GRAY_HEX)}</w:sdtContent></w:sdt>")
    par._p.append(parse_xml(xml))


def dropdown(par, items, hint="選ぶ▼", size=10.5, alias=""):
    if PRINT:
        text(par, "・".join(items[:3]), size)
        return
    li = "".join(f'<w:listItem w:displayText="{escape(i)}" w:value="{escape(i)}"/>' for i in items)
    field(par, hint, size, alias, kind=f"<w:dropDownList>{li}</w:dropDownList>")


def date_field(par, hint="日付を選ぶ▼", size=11, alias="収穫日"):
    if PRINT:
        text(par, "　　月　　日", size)
        return
    kind = ('<w:date><w:dateFormat w:val="M月d日"/><w:lid w:val="ja-JP"/>'
            '<w:storeMappedDataAs w:val="dateTime"/><w:calendar w:val="gregorian"/></w:date>')
    field(par, hint, size, alias, kind=kind)


def checkbox(par, label, size=10):
    if PRINT:
        text(par, "□ " + label, size)
        return
    xml = (f"<w:sdt {nsdecls('w', 'w14')}><w:sdtPr>{_rpr(size, font=BOX_FONT)}"
           f'<w:alias w:val="{escape(label)}"/><w:id w:val="{next(_ids)}"/>'
           '<w14:checkbox><w14:checked w14:val="0"/>'
           f'<w14:checkedState w14:val="2611" w14:font="{BOX_FONT}"/>'
           f'<w14:uncheckedState w14:val="2610" w14:font="{BOX_FONT}"/></w14:checkbox>'
           f"</w:sdtPr><w:sdtContent>{_run('☐', size, font=BOX_FONT)}</w:sdtContent></w:sdt>")
    par._p.append(parse_xml(xml))
    text(par, label, size)


def wrap_block(par, alias, kind, placeholder):
    """段落をブロックレベルのコンテンツ コントロールで包む"""
    plc = "<w:showingPlcHdr/>" if placeholder else ""
    sdt = parse_xml(f"<w:sdt {nsdecls('w')}><w:sdtPr><w:alias w:val=\"{escape(alias)}\"/>"
                    f'<w:id w:val="{next(_ids)}"/>{plc}{kind}</w:sdtPr><w:sdtContent/></w:sdt>')
    par._p.addprevious(sdt)
    sdt.find(qn("w:sdtContent")).append(par._p)


def add_binding_store(doc, values):
    """表紙とヘッダーの「年・組・番・名前」を連動させるためのデータ（customXml）"""
    body = "".join(f"<{k}>{escape(v)}</{k}>" for k, v in values.items())
    item_xml = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><report xmlns="{XML_NS}">{body}</report>'
    props_xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                 f'<ds:datastoreItem ds:itemID="{STORE_ID}" '
                 'xmlns:ds="http://schemas.openxmlformats.org/officeDocument/2006/customXml">'
                 f'<ds:schemaRefs><ds:schemaRef ds:uri="{XML_NS}"/></ds:schemaRefs></ds:datastoreItem>')
    pkg = doc.part.package
    item_name = pkg.next_partname("/customXml/item%d.xml")
    props_name = PackURI(item_name.replace("/item", "/itemProps"))
    item = Part(item_name, CT.XML, item_xml.encode("utf-8"), pkg)
    props = Part(props_name, CT.OFC_CUSTOM_XML_PROPERTIES,
                 props_xml.encode("utf-8"), pkg)
    item.relate_to(props, RT.CUSTOM_XML_PROPS)
    doc.part.relate_to(item, RT.CUSTOM_XML)


# ---------- 部品 ----------

def heading_bar(doc, title, sub=None, page_break=False, size=13):
    """緑の帯の見出し（1行テーブル）。page_break=True でこの見出しから新しいページ。"""
    if page_break:
        p = doc.add_paragraph()
        tight(p)
        p.paragraph_format.page_break_before = True
        p.paragraph_format.line_spacing = Pt(1)
    t = doc.add_table(rows=1, cols=1)
    fix_layout(t, [BODY_W])
    set_cell_margins(t, 1.2, 1.2, 2.5, 2.5)
    c = t.cell(0, 0)
    shade(c, ACCENT)
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    par = write(c, title, size=size, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    if sub:
        par.paragraph_format.tab_stops.add_tab_stop(Mm(BODY_W - 5), WD_TAB_ALIGNMENT.RIGHT)
        text(par, "\t" + sub, 9.5, color=RGBColor(0xFF, 0xFF, 0xFF))
    return t


def label(doc, title, hint=None):
    p = doc.add_paragraph()
    tight(p, before=3, after=1)
    text(p, "■ " + title, 10.5, bold=True, color=RGBColor(0x2F, 0x7D, 0x4F))
    if hint:
        text(p, "　" + hint, 8, color=GRAY)
    return p


def ruled_lines(doc, n, h=7.6):
    """点線の罫線（手書き用）"""
    t = doc.add_table(rows=n, cols=1)
    fix_layout(t, [BODY_W])
    set_cell_margins(t, 0, 0, 1.5, 1.5)
    dotted = ("dotted", 6, "9AA5B1")
    for i, row in enumerate(t.rows):
        row_height(row, h)
        c = row.cells[0]
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
        set_cell_borders(c, top=dotted if i == 0 else None, bottom=dotted, left=None, right=None)
        first_par(c)
    return t


LINE_PITCH = 7.5   # 考察欄の行の高さ（mm）。文字の行送りと罫線の間隔をこれでそろえる
_docpr_ids = iter(range(500000, 10 ** 7))
_ruled_png = {}


def ruled_png(w_mm, n, pitch):
    """行送りと同じ間隔で点線を引いた透明の画像（文字の背面に置く）"""
    key = (w_mm, n, pitch)
    if key not in _ruled_png:
        from PIL import Image, ImageDraw
        k = 10  # px/mm
        W, H = int(w_mm * k), int(n * pitch * k)
        im = Image.new("RGBA", (W, H), (0, 0, 0, 0))
        d = ImageDraw.Draw(im)
        for i in range(1, n + 1):
            y = int(i * pitch * k) - 2
            for x in range(0, W, 12):
                d.line([(x, y), (min(x + 6, W), y)], fill=(150, 162, 178, 255), width=2)
        buf = io.BytesIO()
        im.save(buf, "PNG")
        _ruled_png[key] = buf.getvalue()
    return io.BytesIO(_ruled_png[key])


def behind_image(par, stream, w_mm, h_mm, y_mm=0.0):
    """段落を基準に、文字の背面へ画像を固定配置する（wrapNone / behindDoc）"""
    rid, _ = par.part.get_or_add_image(stream)
    cx, cy, y = int(Mm(w_mm)), int(Mm(h_mm)), int(Mm(y_mm))
    pid = next(_docpr_ids)
    xml = (
        f'<w:r {nsdecls("w", "wp", "a", "pic", "r")}><w:drawing>'
        '<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="1" '
        'behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="column"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="paragraph"><wp:posOffset>{y}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
        f'<wp:docPr id="{pid}" name="罫線{pid}"/>'
        '<wp:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></wp:cNvGraphicFramePr>'
        '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        f'<pic:pic><pic:nvPicPr><pic:cNvPr id="{pid}" name="lines.png"/><pic:cNvPicPr/></pic:nvPicPr>'
        f'<pic:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
        f'<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic>'
        '</a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>')
    par._p.append(parse_xml(xml))


def writing_box(doc, lines, hint, alias, pitch=LINE_PITCH):
    """考察などを書く欄（罫線つき）。

    Word 用：1つの枠の中で自由に書ける入力欄。文字の行送りを罫線の間隔と同じ「固定値」にし、
    罫線は文字の背面に置いた画像なので、打った文字がそのまま罫線の上に並ぶ。
    印刷用：表の罫線。"""
    if PRINT:
        return ruled_lines(doc, lines, h=pitch)
    pad_top, pad_bottom, anchor_pt = 0.6, 0.6, 1
    t = doc.add_table(rows=1, cols=1)
    fix_layout(t, [BODY_W])
    set_cell_margins(t, pad_top, pad_bottom, 2.2, 2.2)
    row = t.rows[0]
    row_height(row, pad_top + pad_bottom + anchor_pt * 0.3528 + lines * pitch + 0.4)
    c = row.cells[0]
    box_borders(c, "9AA5B1", 6)
    # 1段落目：罫線画像の置き場所（高さ 1pt。生徒が触らない位置）
    holder = first_par(c)
    holder.paragraph_format.line_spacing = Pt(anchor_pt)
    behind_image(holder, ruled_png(BODY_W - 4.4, lines, pitch), BODY_W - 4.4, lines * pitch,
                 y_mm=anchor_pt * 0.3528)
    # 2段落目：入力欄。行送りを罫線と同じ固定値に
    par = c.add_paragraph()
    tight(par)
    par.paragraph_format.line_spacing = Mm(pitch)
    text(par, hint, 10.5, color=GRAY)
    wrap_block(par, alias, "<w:richText/>", placeholder=True)
    return t


_placeholders = {}


def placeholder_png(w_mm, h_mm):
    """写真欄に置く「クリックして写真を挿入」の画像"""
    key = (w_mm, h_mm)
    if key in _placeholders:
        return io.BytesIO(_placeholders[key])
    from PIL import Image, ImageDraw, ImageFont
    k = 8
    W, H = int(w_mm * k), int(h_mm * k)
    im = Image.new("RGB", (W, H), (246, 248, 250))
    d = ImageDraw.Draw(im)
    font = None
    for path in _jp_fonts():
        try:
            font = ImageFont.truetype(path, int(5.5 * k))
            small = ImageFont.truetype(path, int(3.4 * k))
            break
        except OSError:
            continue
    cx, cy = W // 2, H // 2
    # カメラの絵
    s = 3.2 * k
    d.rounded_rectangle([cx - 3 * s, cy - 4.2 * s, cx + 3 * s, cy - 0.2 * s], radius=int(s * .6),
                        outline=(150, 160, 175), width=int(k * .6))
    d.ellipse([cx - 1.3 * s, cy - 3.5 * s, cx + 1.3 * s, cy - 0.9 * s],
              outline=(150, 160, 175), width=int(k * .6))
    if font:
        msg, sub = "クリックして写真を挿入", "写真は枠の大きさに合わせて入ります"
        tw = d.textlength(msg, font=font)
        d.text((cx - tw / 2, cy + 1.2 * k), msg, fill=(120, 130, 145), font=font)
        tw = d.textlength(sub, font=small)
        d.text((cx - tw / 2, cy + 9 * k), sub, fill=(150, 160, 175), font=small)
    buf = io.BytesIO()
    im.save(buf, "PNG")
    _placeholders[key] = buf.getvalue()
    return io.BytesIO(_placeholders[key])


def _jp_fonts():
    try:
        out = subprocess.run(["fc-match", "-f", "%{file}", ":lang=ja"], capture_output=True,
                             text=True, timeout=10).stdout.strip()
        if out:
            yield out
    except (OSError, subprocess.SubprocessError):
        pass
    yield from ("C:/Windows/Fonts/BIZ-UDGothicR.ttc", "C:/Windows/Fonts/meiryo.ttc",
                "/System/Library/Fonts/ヒラギノ角ゴシック W3.ttc")


def photo_box(doc, height_mm, image=None, caption="ここに写真を貼る"):
    t = doc.add_table(rows=1, cols=1)
    fix_layout(t, [BODY_W])
    set_cell_margins(t, 1, 1, 1, 1)
    row = t.rows[0]
    row_height(row, height_mm)
    c = row.cells[0]
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    box_borders(c)
    par = first_par(c, WD_ALIGN_PARAGRAPH.CENTER)
    max_w, max_h = BODY_W - 6, height_mm - 5
    if image:
        w, h = fit_image(image, max_w, max_h)
        par.add_run().add_picture(image, width=Mm(w), height=Mm(h))
        if not PRINT:
            wrap_block(par, caption, "<w:picture/>", placeholder=False)
    elif PRINT:
        shade(c, "F6F8FA")
        text(par, caption, 12, color=GRAY)
        p2 = c.add_paragraph()
        tight(p2)
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        text(p2, "（写真をのりで貼る）", 8.5, color=GRAY)
    else:
        par.add_run().add_picture(placeholder_png(max_w, max_h), width=Mm(max_w), height=Mm(max_h))
        wrap_block(par, caption, "<w:picture/>", placeholder=True)
    return t


def fit_image(path, max_w, max_h):
    try:
        from PIL import Image, ImageOps
        with Image.open(path) as im:
            im = ImageOps.exif_transpose(im)
            iw, ih = im.size
    except Exception:
        iw, ih = 4, 3
    scale = min(max_w / iw, max_h / ih)
    return iw * scale, ih * scale


def form_row(doc, cells, widths, h=11):
    """cells: [(label, None) | (None, filler(par)), ...] を1行に並べる記入欄"""
    t = doc.add_table(rows=1, cols=len(cells))
    fix_layout(t, widths)
    set_cell_margins(t)
    row = t.rows[0]
    row_height(row, h)
    for c, (lab, fill) in zip(row.cells, cells):
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        box_borders(c)
        if lab is not None:
            shade(c, PALE)
            write(c, lab, size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        else:
            fill(first_par(c, WD_ALIGN_PARAGRAPH.RIGHT))
    return t


def num(hint, unit, alias):
    """数字の入力欄＋単位"""
    def fill(par):
        field(par, hint, 11, alias)
        text(par, " " + unit, 10.5)
    return fill


# ---------- ページ ----------

def setup(doc, crop, a):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(PAGE_W), Mm(PAGE_H)
    sec.left_margin = sec.right_margin = Mm(MARGIN_X)
    sec.top_margin, sec.bottom_margin = Mm(MARGIN_TOP), Mm(MARGIN_BOTTOM)
    sec.header_distance, sec.footer_distance = Mm(7), Mm(6)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(k), FONT)
    tight(normal)

    # ヘッダー：表紙以外の全ページ。年・組・番・名前は表紙の入力欄と連動
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]
    tight(hp)
    ts = hp.paragraph_format.tab_stops
    for tw in (4680, 9360):  # Header スタイル既定のタブ（中央・右）を打ち消す
        ts.add_tab_stop(Pt(tw / 20), WD_TAB_ALIGNMENT.CLEAR)
    ts.add_tab_stop(Mm(BODY_W), WD_TAB_ALIGNMENT.RIGHT)
    text(hp, f"{crop}の栽培観察レポート\t", 8.5, color=GRAY)
    for key, hint, unit in (("grade", "　", "年 "), ("klass", "　", "組 "),
                            ("number", "　", "番　名前 "), ("name", "＿＿＿＿＿＿", "")):
        field(hp, hint, 8.5, key, value=getattr(a, key), bind=key)
        if unit:
            text(hp, unit, 8.5, color=GRAY)

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(fp)


def fmt_date(d, year=True):
    s = f"{d.month}月{d.day}日（{WEEKDAY[d.weekday()]}）"
    return f"{d.year}年{s}" if year else s


def cover(doc, a, days):
    spacer(doc, 60)
    p = doc.add_paragraph()
    tight(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    text(p, "技術・家庭科（技術分野）　生物育成の技術", 12, color=GRAY)
    spacer(doc, 14)
    p = doc.add_paragraph()
    tight(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    text(p, f"{a.crop}の", 30, bold=True, color=RGBColor(0x2F, 0x7D, 0x4F))
    text(p, "栽培観察レポート", 30, bold=True)
    spacer(doc, 50)

    t = doc.add_table(rows=4, cols=2)
    fix_layout(t, [38, BODY_W - 38])
    set_cell_margins(t, 1, 1, 3, 3)
    period = f"{fmt_date(days[0])}　〜　{fmt_date(days[-1])}"

    def who(par):
        for key, unit in (("grade", " 年　"), ("klass", " 組　"), ("number", " 番")):
            field(par, "数字", 15, key, value=getattr(a, key), bind=key)
            text(par, unit, 15)

    rows = [
        ("年・組・番", who, 20),
        ("名　前", lambda par: field(par, "名前を入力", 15, "名前", value=a.name, bind="name"), 20),
        ("観察・栽培期間", lambda par: text(par, period, 15), 20),
        ("観察日数", lambda par: text(par, f"{len(days)} 日（土日{'・休日' if a.skip else ''}をのぞく）", 12), 14),
    ]
    for row, (lab, fill, h) in zip(t.rows, rows):
        row_height(row, h)
        for c in row.cells:
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            box_borders(c, "5B6573", 10)
        shade(row.cells[0], PALE)
        write(row.cells[0], lab, size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        fill(first_par(row.cells[1], WD_ALIGN_PARAGRAPH.CENTER))

    if not PRINT:
        spacer(doc, 30)
        p = doc.add_paragraph()
        tight(p)
        p.paragraph_format.left_indent = Mm(8)
        p.paragraph_format.line_spacing = 1.4
        text(p, "【記入のしかた】\n", 9.5, bold=True, color=GRAY)
        text(p, "・灰色の文字の欄をクリックして入力します（入力すると灰色の文字は消えます）。\n"
                "・ここで年・組・番・名前を入れると、すべてのページの上に自動で入ります。\n"
                "・写真の枠をクリックすると写真を選べます。写真は枠の大きさに合わせて入ります。\n"
                "・天気は一覧から選びます。手入れは、したものの四角をクリックしてチェックを入れます。\n"
                "・考察の枠に入りきらない分は表示されません。枠に収まるようにまとめましょう。",
             9, color=GRAY)


def day_page(doc, a, n, d, total, image):
    heading_bar(doc, f"第 {n} 日目　　{fmt_date(d, year=False)}",
                sub=f"{d.year}年　（{n} / {total}）", page_break=True, size=14)
    spacer(doc, 3)

    def center(fn):
        def fill(par):
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fn(par)
        return fill

    form_row(doc, [("天気", None), (None, center(lambda p: dropdown(p, WEATHER, alias="天気"))),
                   ("気温", None), (None, num("数字", "℃", "気温")),
                   ("観察時刻", None), (None, center(lambda p: (field(p, "数字", 10.5, "時"), text(p, " 時 "),
                                                               field(p, "数字", 10.5, "分"), text(p, " 分"))))],
             [16, 58, 16, 30, 20, 40], h=9)
    spacer(doc, 3)

    photo_box(doc, 112, image, caption="今日の写真")
    spacer(doc, 3)

    form_row(doc, [("一番大きい株の背丈", None), (None, num("数字", "cm", "背丈")),
                   ("葉の枚数", None), (None, num("数字", "枚", "葉の枚数")),
                   ("前回からの変化", None),
                   (None, lambda p: (field(p, "数字", 11, "背丈の変化"), text(p, " cm／"),
                                     field(p, "数字", 11, "葉の変化"), text(p, " 枚")))],
             [27, 33, 22, 30, 26, 42], h=13)

    label(doc, "今日した手入れ", "（したものに✓）")
    p = doc.add_paragraph()
    tight(p)
    p.paragraph_format.left_indent = Mm(2)
    for item in CARE_ITEMS:
        checkbox(p, item, 9.5)
        text(p, "　", 9.5)
    checkbox(p, "その他（", 9.5)
    field(p, "内容", 9.5, "その他の手入れ")
    text(p, "）", 9.5)

    label(doc, "観察・手入れの記録と考察",
          "様子の変化／何をなぜしたか／変化の理由として考えられること／次にすること")
    writing_box(doc, 12, "ここをクリックして、観察して気づいたこと・した手入れとその理由・考察を書く",
                "観察・手入れの記録と考察")


def summary_page(doc, a, days, image):
    heading_bar(doc, "収穫後のまとめ", sub=f"{a.crop}", page_break=True, size=14)
    spacer(doc, 3)
    form_row(doc, [("収穫日", None), (None, lambda p: date_field(p)),
                   ("収穫量", None),
                   (None, lambda p: (field(p, "数字", 11, "収穫した個数"), text(p, " 個／"),
                                     field(p, "数字", 11, "収穫した重さ"), text(p, " g")))],
             [24, 52, 24, 80], h=11)
    form_row(doc, [("最後の背丈", None), (None, num("数字", "cm", "最後の背丈")),
                   ("葉の枚数", None), (None, num("数字", "枚", "最後の葉の枚数")),
                   ("栽培日数", None),
                   (None, lambda p: text(p, f"{(days[-1] - days[0]).days + 1} 日", 11))],
             [24, 36, 22, 30, 24, 44], h=11)
    spacer(doc, 3)
    photo_box(doc, 70, image, caption="収穫したものの写真")

    for title, hint, n in [
        ("育ててわかったこと・成長の様子", "背丈・葉の数の記録から読み取れる変化など", 5),
        ("工夫したこと・うまくいったこと", "手入れの工夫とその結果", 4),
        ("うまくいかなかったこと・次に育てるときの改善点", "原因と改善の方法", 4),
        ("生物育成の技術と、生活や社会とのつながり", "農家の工夫・食料・環境などと比べて", 4),
    ]:
        label(doc, title, hint)
        writing_box(doc, n, "ここをクリックして書く", title)


# ---------- 日付・写真 ----------

def parse_date(s):
    return dt.date.fromisoformat(s.strip())


def weekdays(start, end, skip):
    out, d = [], start
    while d <= end:
        if d.weekday() < 5 and d not in skip:
            out.append(d)
        d += dt.timedelta(days=1)
    return out


def find_photos(folder):
    """YYYY-MM-DD / YYYYMMDD / MMDD を含むファイル名を日付に対応させる。harvest / 収穫 はまとめ用。"""
    by_date, harvest = {}, None
    if not folder:
        return by_date, harvest
    for name in sorted(os.listdir(folder)):
        if not name.lower().endswith(PHOTO_EXT):
            continue
        path = os.path.join(folder, name)
        stem = os.path.splitext(name)[0]
        if re.search(r"harvest|収穫|まとめ", stem, re.I):
            harvest = path
            continue
        m = re.search(r"(20\d\d)[-_.]?(\d\d)[-_.]?(\d\d)", stem)
        if m:
            key = (int(m[2]), int(m[3]))
        else:
            m = re.search(r"(?<!\d)(\d\d)[-_.]?(\d\d)(?!\d)", stem)
            if not m:
                continue
            key = (int(m[1]), int(m[2]))
        by_date.setdefault(key, path)
    return by_date, harvest


def main():
    global PRINT
    ap = argparse.ArgumentParser(description="栽培観察レポート（Word・A4・1日1枚）を作成")
    ap.add_argument("--crop", required=True, help="作物名（例：ミニトマト）")
    ap.add_argument("--start", required=True, help="観察・栽培の開始日 YYYY-MM-DD")
    ap.add_argument("--end", required=True, help="観察・栽培の終了日 YYYY-MM-DD")
    ap.add_argument("--grade", default="", help="学年")
    ap.add_argument("--klass", default="", help="組")
    ap.add_argument("--number", default="", help="出席番号")
    ap.add_argument("--name", default="", help="名前")
    ap.add_argument("--skip", default="", help="除外する平日（祝日・行事など）カンマ区切り YYYY-MM-DD")
    ap.add_argument("--photos", default="", help="写真フォルダ（ファイル名に日付を含める）")
    ap.add_argument("--print", dest="print_mode", action="store_true",
                    help="手書き用（罫線つき・入力欄なし）で作る")
    ap.add_argument("--pdf", action="store_true", help="LibreOffice で PDF も作る")
    ap.add_argument("-o", "--out", default="", help="出力 .docx（省略時：<作物名>_栽培観察レポート.docx）")
    a = ap.parse_args()
    PRINT = a.print_mode

    start, end = parse_date(a.start), parse_date(a.end)
    if end < start:
        sys.exit("終了日が開始日より前です")
    skip = {parse_date(s) for s in a.skip.split(",") if s.strip()}
    a.skip = skip
    days = weekdays(start, end, skip)
    if not days:
        sys.exit("期間内に平日（月〜金）がありません")
    photos, harvest = find_photos(a.photos)

    doc = Document()
    if not PRINT:
        add_binding_store(doc, {k: getattr(a, k) for k in ("grade", "klass", "number", "name")})
    setup(doc, a.crop, a)
    cover(doc, a, days)
    for i, d in enumerate(days, 1):
        day_page(doc, a, i, d, len(days), photos.get((d.month, d.day)))
    summary_page(doc, a, days, harvest)
    normalize(doc.element)
    normalize(doc.sections[0].header._element)

    suffix = "_手書き用" if PRINT else ""
    out = a.out or f"{a.crop}_栽培観察レポート{suffix}.docx"
    doc.save(out)
    print(f"作成：{out}（{'手書き用' if PRINT else 'Word 入力用'}／表紙1＋観察{len(days)}日＋まとめ1＝{len(days) + 2}ページ）")
    print(f"期間：{fmt_date(days[0])} 〜 {fmt_date(days[-1])}")
    if photos or harvest:
        used = sum(1 for d in days if (d.month, d.day) in photos)
        print(f"写真：{used}枚を日付に対応して貼り付け" + ("、収穫写真あり" if harvest else ""))

    if a.pdf:
        outdir = os.path.dirname(os.path.abspath(out))
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", outdir, out],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        print(f"作成：{os.path.splitext(out)[0]}.pdf")


if __name__ == "__main__":
    main()
